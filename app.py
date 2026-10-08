import io
import os
import zipfile
import pandas as pd
import json
from PIL import Image as PILImage
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import column_index_from_string
import streamlit as st

# Default template filename expected in the repository
DEFAULT_TEMPLATE_PATH = "L-RG-84927-150.xlsx"

# --- Persistent Mapping Storage ---
MAPPINGS_FILE = "mappings.json"

DEFAULT_MAPPINGS = {
    "shapes": [
        {
            "Input Shape (Col D)": "RD(F/C)",
            "Template Value (Col C)": "ROUND BRILLIANT",
        },
        {"Input Shape (Col D)": "OV", "Template Value (Col C)": "OVAL"},
    ],
    "settings": [
        {
            "Input Setting (Col J)": "PRONG",
            "Template Value (Col X)": "Hand Set Prong/Pave/Bead",
        },
        {
            "Input Setting (Col J)": "MICRO SPLIT PRONG",
            "Template Value (Col X)": "Hand Set Micro Split Prong",
        },
    ],
}


def load_persistent_mappings():
    if os.path.exists(MAPPINGS_FILE):
        try:
            with open(MAPPINGS_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return DEFAULT_MAPPINGS
    return DEFAULT_MAPPINGS


def save_persistent_mappings(shapes_data, settings_data):
    with open(MAPPINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(
            {"shapes": shapes_data, "settings": settings_data}, f, indent=4
        )
        
def fit_image_dimensions(image_bytes, max_width=420, max_height=420):
    """Calculates fitted dimensions preserving original aspect ratio."""
    with PILImage.open(io.BytesIO(image_bytes)) as img:
        orig_w, orig_h = img.size
    ratio = min(max_width / orig_w, max_height / orig_h)
    return int(orig_w * ratio), int(orig_h * ratio)


def process_files(
    template_bytes,
    input_bytes,
    shape_mapping,
    setting_mapping,
    cfg,
):
    """Extracts data from the input Excel and populates template sheets in memory."""
    wb_in = openpyxl.load_workbook(io.BytesIO(input_bytes), data_only=True)
    ws_in = wb_in.active

    # 1. Map sketch images anchored in Column O (Col index 14)
    images_by_row = {}
    for img in ws_in._images:
        anchor = img.anchor
        if hasattr(anchor, "_from") and anchor._from.col == 14:
            images_by_row[anchor._from.row] = img._data()

    # 2. Extract style blocks
    styles_data = []
    for r in range(1, ws_in.max_row + 1):
        cell_a = ws_in.cell(row=r, column=1).value
        if cell_a and "Style No." in str(cell_a):
            style_no = str(ws_in.cell(row=r, column=3).value).strip()

            # Locate sketch image
            img_data = None
            for offset in range(-1, 12):
                target_r = (r - 1) + offset
                if target_r in images_by_row:
                    img_data = images_by_row[target_r]
                    break

            # Read diamond details rows
            first_diamond_row = r + 2
            curr_r = first_diamond_row
            diamond_rows = []

            cols_cfg = cfg.get("input_cols", {})
            col_qty_idx = column_index_from_string(cols_cfg.get("qty", "H"))
            col_shape_idx = column_index_from_string(cols_cfg.get("shape", "D"))
            col_sieve_idx = column_index_from_string(cols_cfg.get("sieve", "E"))
            col_mm_idx = column_index_from_string(cols_cfg.get("mm", "F"))
            col_weight_idx = column_index_from_string(cols_cfg.get("weight", "I"))
            col_setting_idx = column_index_from_string(cols_cfg.get("setting", "J"))

            while curr_r <= ws_in.max_row:
                qty_val = ws_in.cell(row=curr_r, column=col_qty_idx).value
                if qty_val is None or str(qty_val).strip() == "":
                    break

                val_d = ws_in.cell(row=curr_r, column=col_shape_idx).value
                val_e = ws_in.cell(row=curr_r, column=col_sieve_idx).value
                val_f = ws_in.cell(row=curr_r, column=col_mm_idx).value
                val_h = qty_val
                val_i = ws_in.cell(row=curr_r, column=col_weight_idx).value
                val_j = ws_in.cell(row=curr_r, column=col_setting_idx).value

                # Parse Sieve Range
                min_sieve, max_sieve = None, None
                if val_e is not None and str(val_e).strip() != "":
                    sieve_str = str(val_e).strip()
                    if "-" in sieve_str:
                        parts = sieve_str.split("-")
                        min_sieve, max_sieve = parts[0].strip(), parts[1].strip()
                    else:
                        min_sieve = sieve_str

                diamond_rows.append(
                    {
                        "shape": str(val_d).strip() if val_d else None,
                        "min_sieve": min_sieve,
                        "max_sieve": max_sieve,
                        "mm": val_f,
                        "qty": val_h,
                        "weight": val_i,
                        "setting": str(val_j).strip() if val_j else None,
                    }
                )
                curr_r += 1

            styles_data.append(
                {
                    "style_no": style_no,
                    "image_bytes": img_data,
                    "diamond_rows": diamond_rows,
                }
            )

    # 3. Populate template workbooks into an in-memory ZIP archive
    zip_buffer = io.BytesIO()
    total_styles = len(styles_data)

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        progress_bar = st.progress(0)
        status_text = st.empty()

        for idx, item in enumerate(styles_data):
            style_no = item["style_no"]
            img_bytes = item["image_bytes"]
            diamond_rows = item["diamond_rows"]

            status_text.text(f"Processing {idx + 1}/{total_styles}: {style_no}")

            wb_tmpl = openpyxl.load_workbook(io.BytesIO(template_bytes))
            ws_tmpl = wb_tmpl["COST SHEET"]

            # Style number placement
            ws_tmpl[cfg["style_cell"]].value = style_no

            # Sketch image placement
            if img_bytes:
                img_w, img_h = fit_image_dimensions(img_bytes)
                xl_img = XLImage(io.BytesIO(img_bytes))
                xl_img.width = img_w
                xl_img.height = img_h
                ws_tmpl.add_image(xl_img, cfg["image_anchor"])

            # Stone type determination
            stone_type = (
                "DIAMOND (LAB)"
                if style_no.upper().startswith("L")
                else "DIAMOND (NAT)"
            )

            # Populate diamond rows
            start_row = cfg["diamond_start_row"]
            for i, d_row in enumerate(diamond_rows):
                tr = start_row + i
                ws_tmpl[f"B{tr}"].value = stone_type
                raw_shape = d_row["shape"]
                ws_tmpl[f"C{tr}"].value = shape_mapping.get(
                    raw_shape, raw_shape
                )
                ws_tmpl[f"G{tr}"].value = "NO"
                ws_tmpl[f"I{tr}"].value = d_row["qty"]
                ws_tmpl[f"K{tr}"].value = d_row["weight"]
                ws_tmpl[f"N{tr}"].value = d_row["mm"]

                min_s = d_row["min_sieve"]
                ws_tmpl[f"P{tr}"].value = (
                    float(min_s)
                    if min_s and min_s.replace(".", "", 1).isdigit()
                    else min_s
                )

                max_s = d_row["max_sieve"]
                ws_tmpl[f"R{tr}"].value = (
                    float(max_s)
                    if max_s and max_s.replace(".", "", 1).isdigit()
                    else max_s
                )

                raw_setting = d_row["setting"]
                ws_tmpl[f"X{tr}"].value = setting_mapping.get(
                    raw_setting, raw_setting
                )

            # Clear leftover template rows (keeps formulas in M, S, V, W intact)
            cols_to_clear = ["B", "C", "E", "G", "I", "K", "N", "P", "R", "X"]
            for empty_r in range(
                start_row + len(diamond_rows), cfg["clear_up_to_row"] + 1
            ):
                for col_letter in cols_to_clear:
                    ws_tmpl[f"{col_letter}{empty_r}"].value = None

            out_buf = io.BytesIO()
            wb_tmpl.save(out_buf)
            zip_file.writestr(f"{style_no}.xlsx", out_buf.getvalue())

            progress_bar.progress((idx + 1) / total_styles)

        status_text.text("Finished processing all styles!")

    zip_buffer.seek(0)
    return zip_buffer, total_styles


# --- Streamlit UI ---
st.set_page_config(page_title="Cost Sheet Generator", layout="wide")
st.title("Excel Cost Sheet Batch Generator")

# Check repository template presence
template_exists = os.path.exists(DEFAULT_TEMPLATE_PATH)


# --- Configuration Section ---
with st.expander("⚙️ View / Modify Mappings & Column Settings", expanded=False):
    current_mappings = load_persistent_mappings()
    tab1, tab2, tab3,tab4 = st.tabs(
        [
            "Shape Mappings",
            "Setting Mappings",
            "Template Layout Coordinates",
            "Input Column Selection"
        ]
    )

    with tab1:
        st.caption(
            "Add, remove, or modify input shape names to their corresponding template dropdown values:"
        )
        shapes_data = current_mappings.get(
            "shapes", DEFAULT_MAPPINGS["shapes"]
        )
        edited_shapes_df = st.data_editor(
            pd.DataFrame(shapes_data), num_rows="dynamic", width="stretch"
        )

    with tab2:
        st.caption(
            "Add, remove, or modify input setting descriptions to their corresponding template dropdown values:"
        )
        settings_data = current_mappings.get(
            "settings", DEFAULT_MAPPINGS["settings"]
        )
        edited_settings_df = st.data_editor(
            pd.DataFrame(settings_data), num_rows="dynamic", width="stretch"
        )

    with tab3:
        st.caption("Customize template target cells and row boundaries:")
        c1, c2, c3, c4 = st.columns(4)
        with c1:
            cfg_style_cell = st.text_input("Style Number Cell", value="D9")
        with c2:
            cfg_img_anchor = st.text_input("Image Top-Left Anchor", value="T7")
        with c3:
            cfg_diamond_start = st.number_input(
                "Diamond Start Row", value=39, step=1
            )
        with c4:
            cfg_clear_up_to = st.number_input(
                "Clear Unused Rows Up To", value=53, step=1
            )
    with tab4:
        st.caption("Select which columns in the input file contain the diamond details:")
        ic1, ic2, ic3 = st.columns(3)
        with ic1:
            in_col_shape = st.text_input("Shape Column", value="D")
            in_col_qty = st.text_input("Qty Column (Row Counter)", value="H")
        with ic2:
            in_col_sieve = st.text_input("Sieve Column", value="E")
            in_col_wt = st.text_input("Total Weight Column", value="I")
        with ic3:
            in_col_mm = st.text_input("MM Column", value="F")
            in_col_setting = st.text_input("Setting Description Column", value="J")
        st.divider()
    if st.button("💾 Save Mappings Permanently"):
        shapes_to_save = edited_shapes_df.dropna(how="all").to_dict(
            orient="records"
        )
        settings_to_save = edited_settings_df.dropna(how="all").to_dict(
            orient="records"
        )
        save_persistent_mappings(shapes_to_save, settings_to_save)
        st.success("Saved successfully! Mappings are now permanently updated.")
# Convert edited DataFrames to lookup dictionaries
shape_mapping = dict(
    zip(
        edited_shapes_df["Input Shape (Col D)"].dropna().str.strip(),
        edited_shapes_df["Template Value (Col C)"].dropna().str.strip(),
    )
)

setting_mapping = dict(
    zip(
        edited_settings_df["Input Setting (Col J)"].dropna().str.strip(),
        edited_settings_df["Template Value (Col X)"].dropna().str.strip(),
    )
)

cfg = {
    "style_cell": cfg_style_cell.strip(),
    "image_anchor": cfg_img_anchor.strip(),
    "diamond_start_row": int(cfg_diamond_start),
    "clear_up_to_row": int(cfg_clear_up_to),
    "input_cols": {
        "shape": in_col_shape.strip().upper(),
        "sieve": in_col_sieve.strip().upper(),
        "mm": in_col_mm.strip().upper(),
        "qty": in_col_qty.strip().upper(),
        "weight": in_col_wt.strip().upper(),
        "setting": in_col_setting.strip().upper(),
    },
}

# --- File Upload & Execution ---
st.divider()
input_upload = st.file_uploader(
    "Upload Input Data File (.xlsx)",
    type=["xlsx"],
    help="Upload the file containing style numbers, sketch images, and diamond rows.",
)

if input_upload:
    if st.button(
        "Generate Cost Sheets", type="primary", disabled=not template_exists
    ):
        with st.spinner("Processing styles and generating Excel sheets..."):
            with open(DEFAULT_TEMPLATE_PATH, "rb") as f:
                template_bytes = f.read()

            input_bytes = input_upload.getvalue()

            zip_result, count = process_files(
                template_bytes=template_bytes,
                input_bytes=input_bytes,
                shape_mapping=shape_mapping,
                setting_mapping=setting_mapping,
                cfg=cfg,
            )

            st.success(f"Successfully generated {count} cost sheets!")
            st.download_button(
                label="📥 Download Generated Files (.zip)",
                data=zip_result,
                file_name="generated_cost_sheets.zip",
                mime="application/zip",
            )
else:
    st.info("Upload your input Excel file above to begin.")