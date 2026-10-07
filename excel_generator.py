import json
import os
from PIL import Image as PILImage
import openpyxl
from openpyxl.drawing.image import Image

TEMPLATE_FILE = "L-RG-84927-150.xlsx"
DATA_JSON = "extracted_data.json"
OUTPUT_DIR = "output"

# Mapping dictionaries
SHAPE_MAPPING = {
    "RD(F/C)": "ROUND BRILLIANT",
    "OV": "OVAL",
}

SETTING_MAPPING = {
    "PRONG": "Hand Set Prong/Pave/Bead",
    "MICRO SPLIT PRONG": "Hand Set Micro Split Prong",
}


def fit_image_to_box(img_path, max_width=450, max_height=420):
    """Calculates resized image dimensions preserving aspect ratio."""
    with PILImage.open(img_path) as pil_img:
        orig_w, orig_h = pil_img.size

    ratio = min(max_width / orig_w, max_height / orig_h)
    new_w = int(orig_w * ratio)
    new_h = int(orig_h * ratio)
    return new_w, new_h


def generate_excels():
    if not os.path.exists(DATA_JSON):
        print(f"Error: {DATA_JSON} not found. Please run extractor.py first.")
        return

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    with open(DATA_JSON, "r", encoding="utf-8") as f:
        styles_data = json.load(f)

    for item in styles_data:
        style_no = item["style_no"]
        image_path = item["image_path"]
        diamond_rows = item["diamond_rows"]

        print(f"Generating Excel for Style: {style_no}...")
        wb = openpyxl.load_workbook(TEMPLATE_FILE)
        ws = wb["COST SHEET"]

        # 1. Fill Style Number in cell D9
        ws["D9"].value = style_no

        # 2. Add Sketch Image to box T7:AB29
        if image_path and os.path.exists(image_path):
            img_w, img_h = fit_image_to_box(
                image_path, max_width=420, max_height=420
            )
            xl_img = Image(image_path)
            xl_img.width = img_w
            xl_img.height = img_h
            ws.add_image(xl_img, "T7")

        # 3. Determine Diamond Type based on style prefix
        stone_type = (
            "DIAMOND (LAB)"
            if style_no.upper().startswith("L")
            else "DIAMOND (NAT)"
        )

        # 4. Fill Diamond Details starting from row 39
        start_row = 39
        row_count = len(diamond_rows)

        for i, d_row in enumerate(diamond_rows):
            target_r = start_row + i

            # Col B: Stone Type
            ws[f"B{target_r}"].value = stone_type

            # Col C: Shape (mapped)
            raw_shape = d_row.get("shape", "")
            ws[f"C{target_r}"].value = SHAPE_MAPPING.get(raw_shape, raw_shape)

            # Col G: Treatment
            ws[f"G{target_r}"].value = "NO"

            # Col I: Qty (from input H)
            ws[f"I{target_r}"].value = d_row.get("qty")

            # Col K: Total Weight (from input I)
            ws[f"K{target_r}"].value = d_row.get("weight")

            # Col N: MM (from input F)
            ws[f"N{target_r}"].value = d_row.get("mm")

            # Col P: Min Sieve (left side of E)
            min_s = d_row.get("min_sieve")
            ws[f"P{target_r}"].value = (
                float(min_s)
                if min_s is not None and min_s.replace(".", "", 1).isdigit()
                else min_s
            )

            # Col R: Max Sieve (right side of E)
            max_s = d_row.get("max_sieve")
            ws[f"R{target_r}"].value = (
                float(max_s)
                if max_s is not None and max_s.replace(".", "", 1).isdigit()
                else max_s
            )

            # Col X: Setting Description (mapped)
            raw_setting = d_row.get("setting", "")
            ws[f"X{target_r}"].value = SETTING_MAPPING.get(
                raw_setting, raw_setting
            )

        # 5. Clear remaining unused diamond rows up to row 53
        # Keeps formulas in M, S, V, W intact
        cols_to_clear = ["B", "C", "E", "G", "I", "K", "N", "P", "R", "X"]
        for empty_r in range(start_row + row_count, 54):
            for col_letter in cols_to_clear:
                ws[f"{col_letter}{empty_r}"].value = None

        # 6. Save the populated workbook as {style_no}.xlsx
        output_filepath = os.path.join(OUTPUT_DIR, f"{style_no}.xlsx")
        wb.save(output_filepath)
        print(f"  Successfully saved to {output_filepath}")

    print(
        f"\nAll {len(styles_data)} files successfully generated in '{OUTPUT_DIR}/'."
    )


if __name__ == "__main__":
    generate_excels()